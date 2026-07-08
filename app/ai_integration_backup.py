"""
Handles sending Task data to AI providers (Anthropic first, OpenAI fallback)
and building a formatted DOCX report.
"""

import os
import re
import requests
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from dotenv import load_dotenv
from .models import Category

load_dotenv()

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")

AI_TIMEOUT = int(os.environ.get("AI_AGENT_TIMEOUT_SECONDS", 60))


# =====================================================
# TASK FORMATTING
# =====================================================

def _tasks_to_text(tasks, include_user=False):
    """Convert tasks to structured text.

    include_user adds a "User:" line per task — used for department-wide
    quarterly reports, where tasks span multiple people, but omitted for
    weekly reports, where every task already belongs to the same author.
    """

    lines = []
    lines.append(f"TOTAL TASKS IN PERIOD: {len(tasks)}")
    lines.append("=" * 80)

    for idx, t in enumerate(tasks, 1):

        lines.append(f"\nTASK #{idx}")
        lines.append(f"Date: {t.date.isoformat()}")

        if include_user:
            lines.append(f"User: {t.owner.name if t.owner else 'Unknown'}")

        lines.append(f"Category: {t.effective_category_name}")
        lines.append(f"Status: {'COMPLETED' if t.is_complete else 'PENDING'}")
        lines.append(f"Description: {t.description}")

        if t.remarks:
            lines.append(f"Remarks: {t.remarks}")

        lines.append("-" * 80)

    return "\n".join(lines)


# =====================================================
# PROMPT BUILDER
# =====================================================

def _build_prompt(tasks_text, report_type, author_name, start_date, end_date):

    category_list = ", ".join(c.name for c in Category.query.order_by(Category.name).all())

    if report_type == "quarterly":

        quarter_number = (start_date.month - 1) // 3 + 1

        prompt = f"""
You are an IT professional writing a comprehensive quarterly report.

DEPARTMENT: {author_name}
PERIOD: {start_date.isoformat()} to {end_date.isoformat()} (Q{quarter_number} {start_date.year})

CRITICAL RULES:

1. You MUST address EVERY task listed. Do not skip or omit any task, even minor ones.
2. Present a holistic view of the quarter — group tasks strictly by the "Category" field given for each task. Use exactly these category names as section headings, in this order, and omit any category with zero tasks: {category_list}.
3. For each task include:
   - sum the tasks up into groups and report on them. DO NOT report them by date. Report on them in paragraphs instead of bullet points.
   - Details of work done, where applicable
   - status
   - remarks if any
   - More IMPORTANTLY, Write the report as paragraphs of full sentences. DO NOT report on the tasks in the groups as bullet points. 
   - Report in clear, often 3rd-person perspective; e.g. 'A meeting with the HoD was held', instead of 'Met with the HoD'
4. Quantify progress wherever possible (e.g. task counts, completion rate per category).
5. Humanize the report as much as possible, but don't introduce grammatical errors. Just ensure that the report does not obviously appear AI-written

REPORT STRUCTURE:

- Executive Summary (headline outcomes across the full quarter)
- Table of Contents
- Introduction
- Key Activities (grouped strictly by Category as instructed above)
- Key Metrics (total tasks, completion rate, breakdown by category)
- Trends and Patterns Observed Across the Quarter
- Challenges and Solutions
- Recommendations for Next Quarter
- Conclusion

TASK DATA:
{tasks_text}

Format the response in markdown.
"""

    else:

        prompt = f"""
You are an IT professional writing a detailed weekly report.

AUTHOR: {author_name}
PERIOD: {start_date.isoformat()} to {end_date.isoformat()}

RULES:

1. Do NOT produce a day-by-day breakdown.
2. Group tasks by the "Category" field provided for each task, not by their free-text Type.
3. Merge descriptions and remarks into human-readable paragraphs.

STRUCTURE:

- Week Overview
- Task Breakdown
- Completed vs Pending Tasks
- Issues Encountered
- Action Items
- Metrics Summary
- Conclusion

TASK DATA:
{tasks_text}

Format the report in markdown.
"""

    return prompt


# =====================================================
# OPENAI CALL
# =====================================================

def _call_openai(prompt, max_tokens=4000, timeout=None):

    headers = {
        "Authorization": f"Bearer {OPENAI_API_KEY}",
        "Content-Type": "application/json",
    }

    body = {
        "model": "gpt-4.1-mini",
        "messages": [
            {
                "role": "system",
                "content": "You are a meticulous IT report writer. Ensure every task is covered.",
            },
            {"role": "user", "content": prompt},
        ],
        "max_tokens": max_tokens,
        "temperature": 0.7,
    }

    resp = requests.post(
        "https://api.openai.com/v1/chat/completions",
        headers=headers,
        json=body,
        timeout=timeout or AI_TIMEOUT,
    )

    resp.raise_for_status()

    data = resp.json()

    return data["choices"][0]["message"]["content"]


# =====================================================
# ANTHROPIC CALL
# =====================================================

def _call_anthropic(prompt, max_tokens=4000, timeout=None):

    headers = {
        "x-api-key": ANTHROPIC_API_KEY,
        "anthropic-version": "2023-06-01",
        "Content-Type": "application/json",
    }

    body = {
        "model": "claude-haiku-4-5-20251001",
        "max_tokens": max_tokens,
        "system": "You are a thorough IT report writer who must include every task.",
        "messages": [
            {"role": "user", "content": prompt}
        ],
    }

    resp = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers=headers,
        json=body,
        timeout=timeout or AI_TIMEOUT,
    )

    resp.raise_for_status()

    data = resp.json()

    # Responses may include non-text blocks (e.g. extended-thinking) ahead of
    # the actual answer, so pick out the text block(s) rather than assuming
    # content[0] is the answer.
    text_blocks = [block["text"] for block in data.get("content", []) if block.get("type") == "text"]

    if not text_blocks:
        raise ValueError(f"Anthropic response contained no text block: {data}")

    return "\n".join(text_blocks)


# =====================================================
# MARKDOWN PARSER
# =====================================================

def _parse_markdown_to_docx(doc, markdown_text):

    lines = markdown_text.split("\n")

    for line in lines:

        stripped = line.strip()

        if not stripped:
            continue

        if stripped.startswith("###"):
            doc.add_heading(stripped.replace("###", "").strip(), level=3)

        elif stripped.startswith("##"):
            doc.add_heading(stripped.replace("##", "").strip(), level=2)

        elif stripped.startswith("#"):
            doc.add_heading(stripped.replace("#", "").strip(), level=1)

        elif stripped.startswith(("- ", "* ", "+ ")):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(stripped[2:])

        elif re.match(r"^\d+\.", stripped):
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\.\s*", "", stripped))

        else:
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.font.size = Pt(11)


# =====================================================
# SHARED REPORT BUILDER
# =====================================================

def _generate_report_bytes(
    tasks,
    report_type_slug,
    report_type_label,
    author_name,
    author_detail,
    start_date,
    end_date,
    include_user_in_tasks=False,
):
    if not tasks:

        doc = Document()

        doc.add_heading(f"{report_type_label} Report - {author_name}", level=1)
        doc.add_paragraph("No tasks found in selected period.")

        bio = BytesIO()
        doc.save(bio)

        return bio.getvalue(), f"report_{report_type_slug}.docx"

    tasks_text = _tasks_to_text(tasks, include_user=include_user_in_tasks)

    prompt = _build_prompt(
        tasks_text,
        report_type_slug,
        author_name,
        start_date,
        end_date,
    )

    final_text = None

    # Quarterly reports must cover far more tasks than weekly ones, so they
    # need a larger output budget and more time to generate.
    if report_type_slug == "quarterly":
        max_tokens = 8000
        ai_timeout = max(AI_TIMEOUT, 120)
    else:
        max_tokens = 4000
        ai_timeout = AI_TIMEOUT

    # =========================
    # TRY ANTHROPIC FIRST
    # =========================

    try:
        if ANTHROPIC_API_KEY:
            final_text = _call_anthropic(prompt, max_tokens=max_tokens, timeout=ai_timeout)
    except Exception as e:
        print(f"Anthropic failed: {e}")

    # =========================
    # FALLBACK TO OPENAI
    # =========================

    if not final_text:

        try:
            if OPENAI_API_KEY:
                final_text = _call_openai(prompt, max_tokens=max_tokens, timeout=ai_timeout)

        except Exception as e:
            print(f"OpenAI failed: {e}")

    if not final_text:
        final_text = "Error: AI providers failed to generate report."

    # =========================
    # BUILD DOCX
    # =========================

    doc = Document()

    title = doc.add_heading(f"{report_type_label} Report", level=1)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    meta = doc.add_paragraph()

    meta.add_run("Author: ").bold = True
    meta.add_run(f"{author_name} ({author_detail})\n")

    meta.add_run("Period: ").bold = True
    meta.add_run(f"{start_date} to {end_date}\n")

    meta.add_run("Total Tasks: ").bold = True
    meta.add_run(str(len(tasks)) + "\n")

    completed = sum(1 for t in tasks if t.is_complete)

    meta.add_run("Completed: ").bold = True
    meta.add_run(f"{completed} | ")

    meta.add_run("Pending: ").bold = True
    meta.add_run(str(len(tasks) - completed))

    doc.add_paragraph()

    _parse_markdown_to_docx(doc, final_text)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    filename = f"report_{report_type_slug}_{start_date}_{end_date}.docx"

    return bio.getvalue(), filename


# =====================================================
# PUBLIC REPORT GENERATORS
# =====================================================

def generate_weekly_report_docs(tasks, user, start_date, end_date):
    """Personal weekly report — tasks belong to a single user."""
    return _generate_report_bytes(
        tasks,
        "weekly",
        "Weekly",
        user.name,
        user.email,
        start_date,
        end_date,
        include_user_in_tasks=False,
    )


def generate_quarterly_report_docs(tasks, department, start_date, end_date):
    """Department-wide quarterly report — tasks span every user in the department."""
    contributor_count = len({t.user_id for t in tasks}) if tasks else 0
    detail = f"{contributor_count} contributor{'s' if contributor_count != 1 else ''}"
    return _generate_report_bytes(
        tasks,
        "quarterly",
        "Quarterly",
        department.name,
        detail,
        start_date,
        end_date,
        include_user_in_tasks=True,
    )